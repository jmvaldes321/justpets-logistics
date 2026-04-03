"""
Descarga el reporte consolidado de inventario desde logystix.co,
procesa el Excel y actualiza api/data.json con los nuevos stocks.
"""
import asyncio
import os
import json
import glob
import sys
import tempfile
from pathlib import Path
from datetime import datetime

import openpyxl
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout


BASE_DIR = Path(__file__).parent.parent
DATA_PATH = BASE_DIR / "api" / "data.json"

USERNAME = os.environ["LOGYSTIX_USER"]
PASSWORD = os.environ["LOGYSTIX_PASSWORD"]


async def download_report(download_dir: str) -> str:
    """Navega logystix.co, genera el reporte y descarga el Excel. Retorna la ruta del archivo."""
    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-setuid-sandbox", "--disable-blink-features=AutomationControlled"]
        )
        context = await browser.new_context(
            accept_downloads=True,
            user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            viewport={"width": 1280, "height": 800}
        )
        page = await context.new_page()

        print("→ Abriendo login...")
        await page.goto("https://ok.logystix.co/site/login", wait_until="domcontentloaded")
        await page.wait_for_timeout(2000)
        await page.screenshot(path="step1_login_page.png")
        print(f"  URL actual: {page.url}")

        # Login
        await page.fill("#username-input", USERNAME)
        await page.fill("#password-input", PASSWORD)
        await page.screenshot(path="step2_filled.png")

        await page.locator("button[type='submit'], .btn-login, input[type='submit'], button:has-text('Entrar')").first.click()
        # Esperar a que el SPA cargue el dashboard completamente
        await page.wait_for_load_state("networkidle", timeout=30000)
        await page.wait_for_timeout(3000)
        await page.screenshot(path="step3_after_login.png")
        print(f"  URL post-login: {page.url}")

        # Detectar login exitoso: no hay formulario de login visible
        login_form_visible = await page.locator("#username-input").is_visible()
        if login_form_visible:
            raise RuntimeError("Login fallido — el formulario sigue visible. Ver step3_after_login.png")
        print("✓ Login exitoso")

        # Esperar a que el sidebar cargue completamente
        await page.wait_for_timeout(4000)
        await page.screenshot(path="step3b_dashboard_loaded.png")

        # Dump de todos los textos visibles del sidebar para debugging
        all_links = await page.locator("a, [role='menuitem'], nav span, aside span, .menu-item, li a").all_text_contents()
        menu_texts = [t.strip() for t in all_links if t.strip()]
        print(f"  Textos del menú: {menu_texts[:40]}")

        # Navegar a Reporte de Inventarios → Reporte Consolidado
        print("→ Navegando a Reporte consolidado de inventarios...")
        # Probar múltiples variantes del texto del menú
        nav1_options = ["Reportes de Inventario", "Reporte de Inventario", "Inventario"]
        nav1_clicked = False
        for text in nav1_options:
            try:
                locator = page.get_by_text(text, exact=False).first
                if await locator.is_visible():
                    await locator.click(timeout=5000)
                    await page.wait_for_load_state("networkidle")
                    await page.wait_for_timeout(1500)
                    nav1_clicked = True
                    print(f"✓ Clic en '{text}'")
                    break
            except Exception:
                continue

        if not nav1_clicked:
            await page.screenshot(path="nav_error.png")
            raise RuntimeError(f"No se encontró ninguna opción de menú de inventario. Textos disponibles: {menu_texts[:30]}")

        await page.screenshot(path="step4_menu_inventario.png")

        try:
            await page.get_by_text("Reporte consolidado", exact=False).first.click(timeout=10000)
            await page.wait_for_load_state("networkidle")
            await page.wait_for_timeout(1500)
        except PlaywrightTimeout:
            # Tomar screenshot y mostrar textos disponibles
            all_texts = await page.locator("a, button, li").all_text_contents()
            available = [t.strip() for t in all_texts if t.strip()]
            await page.screenshot(path="nav2_error.png")
            raise RuntimeError(f"No se encontró 'Reporte consolidado'. Textos disponibles: {available[:30]}")

        print("✓ En página de reporte consolidado")
        await page.screenshot(path="debug_reporte.png")

        # Seleccionar bodega "ebox layout"
        # El campo usa Select2 (data-select2-id) — el native <select> está oculto,
        # hay que interactuar con la UI de Select2.
        print("→ Seleccionando bodega 'ebox layout'...")
        bodega_seleccionada = False

        # 1. Forzar selección en el native <select> con JS y disparar evento change de Select2
        try:
            selected = await page.evaluate("""() => {
                const selects = document.querySelectorAll('select');
                for (const sel of selects) {
                    for (const opt of sel.options) {
                        if (opt.text.toLowerCase().includes('ebox')) {
                            sel.value = opt.value;
                            sel.dispatchEvent(new Event('change', { bubbles: true }));
                            // Trigger Select2 change event if available
                            if (window.jQuery && window.jQuery(sel).data('select2')) {
                                window.jQuery(sel).trigger('change');
                            }
                            return opt.text;
                        }
                    }
                }
                return null;
            }""")
            if selected:
                bodega_seleccionada = True
                print(f"✓ Bodega (JS + Select2): {selected.strip()}")
                await page.wait_for_timeout(500)
        except Exception as e:
            print(f"  JS select fallido: {e}")

        # 2. Si JS no funcionó, interactuar con el UI de Select2
        if not bodega_seleccionada:
            try:
                # Abrir el dropdown de Select2 (contenedor visible)
                await page.locator(".select2-selection, .select2-container").first.click()
                await page.wait_for_timeout(600)
                # Buscar el campo de búsqueda dentro del dropdown abierto
                search_input = page.locator(".select2-search__field")
                if await search_input.is_visible():
                    await search_input.fill("ebox")
                    await page.wait_for_timeout(400)
                # Hacer click en la opción visible
                await page.locator(".select2-results__option:has-text('ebox')").first.click(timeout=5000)
                bodega_seleccionada = True
                print("✓ Bodega (Select2 UI): ebox layout")
            except Exception as e:
                print(f"  Select2 UI fallido: {e}")

        if not bodega_seleccionada:
            await page.screenshot(path="bodega_error.png")
            raise RuntimeError("No se pudo seleccionar la bodega 'ebox layout'. Ver bodega_error.png")

        await page.screenshot(path="step5_bodega_selected.png")

        # Generar reporte
        print("→ Generando reporte...")
        try:
            await page.get_by_text("Generar Reporte", exact=False).first.click()
            await page.wait_for_load_state("networkidle", timeout=60000)
            await page.wait_for_timeout(3000)
        except PlaywrightTimeout:
            await page.screenshot(path="generar_error.png")
            raise RuntimeError("Timeout al generar reporte. Ver generar_error.png")
        await page.screenshot(path="step6_reporte_generado.png")
        print("✓ Reporte generado")

        # Descargar Excel
        # La tabla "Descarga de Reportes" muestra los reportes generados.
        # El más reciente (primer fila) tiene el botón de descarga en la columna ACCIONES.
        print("→ Iniciando descarga del Excel...")
        await page.screenshot(path="step7_before_download.png")

        async with page.expect_download(timeout=90000) as dl_info:
            download_triggered = False
            # 1. Botón de la primera fila en la tabla de reportes (más reciente)
            try:
                first_row_btn = page.locator("table tbody tr:first-child td:last-child a, table tbody tr:first-child td:last-child button").first
                if await first_row_btn.is_visible(timeout=3000):
                    await first_row_btn.click()
                    download_triggered = True
                    print("✓ Click en botón de la primera fila de la tabla")
            except Exception as e:
                print(f"  Primer fila fallida: {e}")

            # 2. Buscar cualquier link/botón de descarga por atributo href o texto
            if not download_triggered:
                try:
                    await page.locator("a[href*='download'], a[href*='excel'], a[download]").first.click(timeout=5000)
                    download_triggered = True
                    print("✓ Click en link de descarga")
                except Exception:
                    pass

            # 3. Buscar botón naranja/acción en la tabla
            if not download_triggered:
                try:
                    await page.locator("table .btn, table button, table a").first.click(timeout=5000)
                    download_triggered = True
                    print("✓ Click en botón de tabla")
                except Exception:
                    pass

            if not download_triggered:
                await page.screenshot(path="download_error.png")
                all_texts = await page.locator("table a, table button").all_text_contents()
                raise RuntimeError(f"No se encontró botón de descarga. Botones en tabla: {all_texts[:10]}")

        download = await dl_info.value
        dest = os.path.join(download_dir, download.suggested_filename or "inventario.xlsx")
        await download.save_as(dest)
        print(f"✓ Descargado: {dest}")

        await browser.close()
        return dest


def parse_inventory_excel(path: str) -> dict:
    """
    Lee el Excel descargado y retorna un dict {ean: cantidad}.
    Busca columnas que contengan SKU/EAN y cantidad/stock/inventario.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    headers = []
    sku_col = None
    qty_col = None

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if any(row):
            headers = [str(c).strip().lower() if c else "" for c in row]
            break

    sku_keywords = ["ean", "sku", "codigo", "código", "barcode", "cod"]
    qty_keywords = ["cantidad", "stock", "inventario", "qty", "unidades", "disponible", "saldo"]

    for idx, h in enumerate(headers):
        if sku_col is None and any(k in h for k in sku_keywords):
            sku_col = idx
        if qty_col is None and any(k in h for k in qty_keywords):
            qty_col = idx

    if sku_col is None or qty_col is None:
        print(f"Headers encontrados: {headers}")
        raise RuntimeError(
            f"No se encontraron columnas de SKU/EAN o cantidad. Headers: {headers}"
        )

    print(f"→ Columna SKU: '{headers[sku_col]}' (col {sku_col}), Cantidad: '{headers[qty_col]}' (col {qty_col})")

    inventory = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        sku = row[sku_col]
        qty = row[qty_col]
        if sku is None:
            continue
        sku_str = str(sku).strip().split(".")[0]  # Eliminar decimales si EAN viene como float
        qty_int = int(qty) if isinstance(qty, (int, float)) else 0
        inventory[sku_str] = qty_int

    print(f"✓ {len(inventory)} SKUs parseados del Excel")
    return inventory


def update_data_json(inventory: dict):
    """Actualiza api/data.json con los nuevos stocks y recalcula M3 totales."""
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        products = json.load(f)

    updated = 0
    not_found = 0
    for p in products:
        sku = p["sku"]
        if sku in inventory:
            old_inv = p["inventario"]
            new_inv = inventory[sku]
            p["inventario"] = new_inv
            p["m3_totales"] = round(p["m3_producto"] * new_inv, 6)
            if old_inv != new_inv:
                updated += 1
        else:
            not_found += 1

    with open(DATA_PATH, "w", encoding="utf-8") as f:
        json.dump(products, f, ensure_ascii=False)

    total_m3 = sum(p["m3_totales"] for p in products)
    print(f"✓ data.json actualizado: {updated} productos modificados, {not_found} no encontrados en el reporte")
    print(f"✓ M³ total bodega: {total_m3:.4f}")


async def main():
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"\n{'='*50}")
    print(f"Sync inventario — {ts}")
    print(f"{'='*50}\n")

    with tempfile.TemporaryDirectory() as tmp:
        excel_path = await download_report(tmp)
        inventory = parse_inventory_excel(excel_path)
        update_data_json(inventory)

    print("\n✓ Sync completo")


if __name__ == "__main__":
    asyncio.run(main())
